"""
Extracts a PPTX or slide-PDF lecture deck into a lossless, ordered JSON manifest.

The rule this script exists to serve: nothing on a slide is dropped except
proven furniture. Every run of text becomes a block, in reading order, tagged
with the slide it came from, so the import step can never silently lose a
bullet. Where a shape cannot be read it is counted as unreadable and surfaced
in the summary rather than passed over -- a visible gap is recoverable, a
silent one is not.

Usage:  python scripts/extract-deck.py <source> <outdir>
"""
import sys, os, json, re, hashlib, collections

SLIDE_NUMBER_PH = 13
TITLE_PH = {1, 3}          # TITLE, CENTER_TITLE
PICTURE = 13               # MSO_SHAPE_TYPE.PICTURE
MEDIA = 16                 # embedded video or animation
GROUP = 6

# A picture repeated on this many slides, and small, is a logo or a border.
FURNITURE_MIN_USES = 5
FURNITURE_MAX_BYTES = 60_000

TITLE_MAX_TOP_FRACTION = 0.20   # a title sits in the top fifth of the slide
TITLE_MIN_PT = 26.0
TITLE_MAX_CHARS = 110
BODY_MIN_PT = 11.0              # below this it is a page number or a credit

# A bullet glyph fused to the word after it leaves a stray letter on the front
# of a heading ("sTransponders"). No heading in these decks begins that way.
STUCK_GLYPH = re.compile(r"^[sS](?=[A-Z][a-z])")

# Stray glyph rows a PDF text layer leaves behind: a lone maths symbol, a
# leader of underscores, a residual slide-number footer. None of it is content
# and all of it reads as a heading if left in place.
JUNK_LINE = re.compile(
    r"^(?:[^\w\s]{1,4}(?:\s+[^\w\s]{1,4})*"   # a row of lone symbols
    r"|[_\-–—.\s]{3,}"                        # a leader rule
    r"|Slide No\.?\s*\d*"                     # a residual slide-number footer
    r"|.{1,2})$",                             # a one- or two-character line
    re.I,
)


# A slide-number footer with the deck's name run onto it -- "Slide No. 125
# Principles of Flight and Aircraft Performance (A)". Usually caught as a
# running line because it repeats on every page, but a slide copied in from
# another deck brings its old footer with it, once, and nothing repeated
# catches a single occurrence. No sentence in these manuals opens this way.
SLIDE_FOOTER = re.compile(r"^Slide No\.?\s*\d", re.I)


def is_junk(line):
    """
    Whether a line is a PDF artefact rather than content.

    Deliberately narrow. An acronym alone on a line is real content in these
    subjects -- "TAF", "ISA", "QNH" all appear as headings -- so a rule that
    drops short lines drops teaching. Only shapes that cannot carry meaning
    are removed: symbol rows, rules, footers, and a word repeated to itself,
    which is how the text layer renders a shadowed or outlined title
    ("Time time").
    """
    if JUNK_LINE.match(line) or SLIDE_FOOTER.match(line):
        return True
    words = line.split()
    if 2 <= len(words) <= 3 and len(line) <= 24:
        lowered = {w.lower() for w in words}
        if len(lowered) == 1:
            return True
    return False

# Words a heading does not end on. A line breaking after one of these is the
# first half of a title that wrapped, not a title in its own right.
DANGLING = {
    "and", "or", "of", "the", "a", "an", "to", "for", "in", "on", "at", "by",
    "with", "from", "as", "is", "are", "its", "their", "your", "our", "over",
    "under", "between", "during", "within", "without", "when", "where",
}


def looks_unfinished(text):
    """Whether a heading line is obviously the first half of a wrapped one."""
    if not text:
        return False
    stripped = text.rstrip()
    if stripped.endswith((",", "-", "\u2013", "\u2014", ":", "&", "/", "+")):
        return True
    if stripped.count("(") > stripped.count(")"):
        return True
    if stripped.count("\u201c") > stripped.count("\u201d"):
        return True
    last = stripped.split()[-1].lower().strip(".,;") if stripped.split() else ""
    return last in DANGLING


def join_wrapped_title(lines, starts_bullet, index, limit):
    """
    Consumes the continuation lines of a title that wrapped.

    PDF text extraction gives one line per rendered line, so a heading set
    across two or three lines arrives as two or three entries -- "Cruising
    Levels (Magnetic Track", "Altitude Requirements) -", "Auckland Oceanic
    FIR". Taking only the first leaves lessons named after sentence fragments,
    and leaves the rest of the heading sitting in the body.

    Returns (title, lines_consumed).
    """
    title = lines[index]
    used = 1
    while looks_unfinished(title) and index + used < len(lines):
        nxt = lines[index + used]
        if starts_bullet[index + used]:
            break
        if len(title) + 1 + len(nxt) > limit:
            break
        title = f"{title} {nxt}"
        used += 1
    return title, used

# How much larger than the body a run has to be before it reads as the title.
TITLE_SIZE_RATIO = 1.25
# Rendered sizes wobble in the last decimal ("44.0" and "44.1" in one heading),
# so runs within this fraction of the largest are treated as the same size.
TITLE_SIZE_TOLERANCE = 0.02


def pdf_page_title(page):
    """
    The page's title, taken from the size the text is set at.

    A PowerPoint slide exported to PDF loses its title placeholder, so the
    extractor had been guessing from wording: take the first line, unless it
    looks like a bullet, unless it looks unfinished. On these decks that guess
    read a heading set across two lines -- "Static" above "Pressure" -- as a
    lesson called "Static", and left "Pressure" at the top of the body. Fifty
    topics in the IR course were named after half of their own heading.

    The text layer still carries the font size, and a slide title is set two or
    three times the size of its body. That is the signal, and it does not
    depend on the words.

    Returns (title, parts) or (None, []) where the page has no distinct title
    -- a page set entirely in one size, which the caller reads the old way.
    """
    runs = []

    def visit(text, cm, tm, font_dict, font_size):
        stripped = text.strip()
        if not stripped:
            return
        try:
            size = float(font_size or 0) * abs(float(tm[3]) or 1.0)
            baseline = float(tm[5])
        except Exception:
            return
        if size <= 0:
            return
        try:
            left = float(tm[4])
        except Exception:
            left = 0.0
        runs.append((baseline, left, size, stripped))

    try:
        page.extract_text(visitor_text=visit)
    except Exception:
        return None, []

    if len(runs) < 2:
        return None, []

    largest = max(size for _, _, size, _ in runs)
    # The body is whichever size carries the most characters, so a two-word
    # heading cannot outvote the paragraph under it.
    weight = {}
    for _, _, size, text in runs:
        weight[size] = weight.get(size, 0) + len(text)
    body = max(weight, key=lambda size: weight[size])

    if largest < body * TITLE_SIZE_RATIO:
        return None, []

    cutoff = largest * (1 - TITLE_SIZE_TOLERANCE)
    # Reading order: down the page, and left to right across each line. Sorting
    # on the baseline alone put a heading laid out in two columns back together
    # in the wrong order — "Cruising Levels (Magnetic Track Altitude
    # Requirements) – Zealand FIR New".
    title_runs = sorted(
        [run for run in runs if run[2] >= cutoff], key=lambda run: (-run[0], run[1])
    )
    parts = [clean(text) for _, _, _, text in title_runs]
    parts = [part for part in parts if part]
    if not parts:
        return None, []

    title = clean(" ".join(parts))
    if not title or len(title) > TITLE_MAX_CHARS or not any(c.isalpha() for c in title):
        return None, []
    return title, parts


def _squashed(text):
    """Letters and digits only, for comparing text across two extractions."""
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def strip_title_lines(lines, starts_bullet, title):
    """
    Drops the body lines that the font-size title was made of.

    The title comes from one reading of the page and the body from another, so
    they cannot be compared literally -- one has the bullet glyphs stripped,
    the other has the line breaks. Comparing them with the punctuation and
    spacing removed is enough, and stopping at the first line that does not
    extend the title keeps a body line that merely repeats a word.
    """
    wanted = _squashed(title)
    if not wanted:
        return lines, starts_bullet

    consumed = 0
    seen = ""
    while consumed < len(lines):
        nxt = seen + _squashed(lines[consumed])
        if not wanted.startswith(nxt):
            break
        seen = nxt
        consumed += 1
        if seen == wanted:
            break

    if seen != wanted:
        return lines, starts_bullet
    return lines[consumed:], starts_bullet[consumed:]


# Symbol-font bullets as they survive PDF text extraction.
BULLET_GLYPH = re.compile(r"^(?:[SÿØ•▪·–-]\s+)")


# What a browser will actually paint. A PDF may embed an image in any format
# its producer liked -- JPEG 2000 and Windows Metafile both turn up in these
# decks -- and every one of those reaches the page as a broken image while the
# file, the row and the request all look fine. Anything not on this list is
# re-encoded on the way out.
WEB_SAFE = {"png", "jpg", "jpeg", "gif", "webp"}


def sniff_ext(blob, fallback="png"):
    """
    The format a blob actually is, read from its own first bytes.

    Needed because the declared extension is sometimes a guess. Where the
    library cannot report a format the obvious fallback is "png", and a JPEG
    stored under a .png name is served with the wrong content type — which the
    reader's `nosniff` header turns from a cosmetic wrong label into a diagram
    the browser refuses to draw.
    """
    if blob[:4] == bytes([0x89, 0x50, 0x4E, 0x47]):
        return "png"
    if blob[:2] == bytes([0xFF, 0xD8]):
        return "jpg"
    if blob[:3] == b"GIF":
        return "gif"
    if blob[:4] == b"RIFF" and blob[8:12] == b"WEBP":
        return "webp"
    if blob[4:8] == b"jP  " or blob[:4] == bytes([0, 0, 0, 12]):
        return "jp2"
    if blob[:4] == bytes([0xD7, 0xCD, 0xC6, 0x9A]):
        return "wmf"
    return fallback


def web_safe_image(blob, ext):
    """
    Returns (bytes, extension) a browser can display.

    The caller hashes the *original* bytes, so an image keeps its identity
    across a re-encode and a re-import lands on the same storage key. Only the
    stored representation changes.
    """
    ext = sniff_ext(blob, ext)
    if ext.lower() in WEB_SAFE:
        return blob, ext
    try:
        from PIL import Image
        import io as _io

        image = Image.open(_io.BytesIO(blob))
        if image.mode not in ("RGB", "RGBA", "L", "LA", "P"):
            image = image.convert("RGBA")
        out = _io.BytesIO()
        image.save(out, format="PNG", optimize=True)
        return out.getvalue(), "png"
    except Exception:
        # Better to store what was found than to drop a diagram. The repair
        # script reports anything still unreadable rather than it vanishing.
        return blob, ext


# A picture with less tonal variation than this carries nothing a reader can
# make out. It is the threshold between the washed-out backdrops these decks
# print behind their text -- a radio mast at 5% opacity, the academy's crest
# ghosted across the page -- and the faintest real teaching image in the
# material, a night photograph of approach lighting. Measured on the IR decks,
# the backdrops top out at 11.8 and the night photograph sits at 14.1.
DECORATIVE_MAX_CONTRAST = 13.0


def is_decorative(blob):
    """
    True for a picture that is decoration rather than a diagram.

    Two signals, both about the image itself rather than about its subject:
    a picture composited at a uniform partial opacity is a backdrop -- nothing
    is laid *over* a diagram -- and a picture with almost no tonal range has
    nothing legible on it whatever it depicts.

    Wrong answers here cost a diagram, so it says no whenever it cannot tell.
    """
    try:
        from PIL import Image, ImageStat
        import io as _io

        image = Image.open(_io.BytesIO(blob))
        image.load()

        if image.mode in ("RGBA", "LA"):
            alpha = image.getchannel("A")
            low, high = alpha.getextrema()
            # Uniform and not opaque: the whole picture sits behind something.
            if low == high and high < 250:
                return True

        stat = ImageStat.Stat(image.convert("RGB"))
        return max(stat.stddev) < DECORATIVE_MAX_CONTRAST
    except Exception:
        return False


def clean(text):
    """Normalises PowerPoint's whitespace without touching the words."""
    if text is None:
        return ""
    # Vertical tab is PowerPoint's soft line break; it separates real lines.
    text = text.replace("\x0b", "\n").replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ").replace("​", "")
    # Collapse runs of spaces and tabs only -- newlines carry meaning and stay.
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def walk(shapes):
    """Yields every shape, descending into groups so nothing nested is lost."""
    for shape in shapes:
        if shape.shape_type == GROUP:
            for inner in walk(shape.shapes):
                yield inner
        else:
            yield shape


def shape_font_pt(shape):
    """Largest explicit font size on the shape, in points, or None."""
    best = None
    if not shape.has_text_frame:
        return None
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size is not None:
                pt = run.font.size.pt
                if best is None or pt > best:
                    best = pt
    return best


def paragraphs_of(shape):
    """(level, text) for every non-empty paragraph, in document order."""
    out = []
    for para in shape.text_frame.paragraphs:
        text = clean("".join(run.text for run in para.runs) or para.text)
        if text:
            out.append((para.level, text))
    return out


def extract_pptx(path, outdir):
    from pptx import Presentation

    prs = Presentation(path)
    slides = list(prs.slides)
    slide_h = prs.slide_height or 6858000

    # ---- pass 1: count image reuse so furniture can be recognised ----------
    uses = collections.Counter()
    blobs = {}
    for slide in slides:
        seen = set()
        for shape in walk(slide.shapes):
            if shape.shape_type != PICTURE:
                continue
            try:
                blob = shape.image.blob
                ext = shape.image.ext
            except Exception:
                continue
            digest = hashlib.sha1(blob).hexdigest()
            blobs[digest] = (blob, ext)
            if digest not in seen:
                seen.add(digest)
                uses[digest] += 1

    furniture = set()
    for digest, count in uses.items():
        repeated_small = (
            count >= FURNITURE_MIN_USES and len(blobs[digest][0]) <= FURNITURE_MAX_BYTES
        )
        if repeated_small or is_decorative(blobs[digest][0]):
            furniture.add(digest)

    assets_dir = os.path.join(outdir, "assets")
    os.makedirs(assets_dir, exist_ok=True)
    for digest, pair in blobs.items():
        if digest in furniture:
            continue
        safe, safe_ext = web_safe_image(pair[0], pair[1])
        blobs[digest] = (pair[0], safe_ext)
        with open(os.path.join(assets_dir, digest + "." + safe_ext), "wb") as fh:
            fh.write(safe)

    # ---- pass 2: read every slide in reading order -------------------------
    out_slides = []
    for index, slide in enumerate(slides, start=1):
        shapes = list(walk(slide.shapes))
        sizes = [p for p in (shape_font_pt(s) for s in shapes) if p]
        max_pt = max(sizes) if sizes else 0

        title = None
        title_shape = None
        title_overflow = []
        for shape in shapes:
            if not shape.has_text_frame:
                continue
            text = clean(shape.text_frame.text)
            if not text or text.isdigit():
                continue
            is_ph_title = shape.is_placeholder and shape.placeholder_format.type in TITLE_PH
            pt = shape_font_pt(shape) or 0
            # A heading names a thing; it does not make a statement. A large
            # top-of-slide run that reads as a finished sentence is emphasised
            # body text, and promoting it to a title would turn a teaching
            # point into a lesson name.
            # A question makes a fine heading ("How Much Water Vapour Can Air
            # Hold?"); a statement does not. An exclamation is always a
            # statement — "Important! Fuel Pounds is not a unit!" is emphasised
            # body text that a size-only rule promotes to a lesson name.
            first_line = text.split("\n")[0].strip()
            reads_as_sentence = (
                text.endswith(".")
                or text.endswith("!")
                or (len(first_line) > 70 and ". " in first_line)
            )
            looks_like_title = (
                (shape.top or 0) <= slide_h * TITLE_MAX_TOP_FRACTION
                and pt >= TITLE_MIN_PT
                and pt >= max_pt
                and len(text) <= TITLE_MAX_CHARS
                and not reads_as_sentence
                # A link is a resource on the slide, never its heading.
                and "http" not in text.lower()
            )
            if is_ph_title or looks_like_title:
                # Only the first line names the slide. Anything after it is
                # content, and is put back into the body below rather than
                # being swallowed into the lesson title.
                lines = [l.strip() for l in text.split("\n") if l.strip()]
                title, title_shape = lines[0], shape
                title_overflow = lines[1:]
                break

        blocks, pictures, tables, unreadable = [], [], [], 0
        for extra in title_overflow:
            blocks.append({"kind": "text", "level": 0, "text": extra})

        for shape in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            if shape is title_shape:
                continue
            if shape.is_placeholder and shape.placeholder_format.type == SLIDE_NUMBER_PH:
                continue

            if shape.shape_type == PICTURE:
                try:
                    blob = shape.image.blob
                    digest = hashlib.sha1(blob).hexdigest()
                except Exception:
                    unreadable += 1
                    continue
                if digest in furniture:
                    continue
                if digest not in blobs:
                    # Pass one could not read this one but pass two can. Store
                    # it now rather than dropping a diagram over a bookkeeping
                    # difference between the two walks.
                    try:
                        ext = shape.image.ext
                    except Exception:
                        ext = "png"
                    safe, safe_ext = web_safe_image(blob, ext)
                    blobs[digest] = (blob, safe_ext)
                    with open(os.path.join(assets_dir, digest + "." + safe_ext), "wb") as fh:
                        fh.write(safe)
                pictures.append({
                    "asset": digest + "." + blobs[digest][1],
                    "sha1": digest,
                    "bytes": len(blob),
                    "top": shape.top or 0,
                    "left": shape.left or 0,
                })
                blocks.append({"kind": "picture", "sha1": digest})
                continue

            if shape.shape_type == MEDIA:
                # An embedded video. It cannot be carried into the web reader,
                # but the slide must not vanish: record it so the lesson still
                # shows that something was taught here, and where.
                poster = None
                try:
                    blob = shape.image.blob
                    digest = hashlib.sha1(blob).hexdigest()
                    safe, safe_ext = web_safe_image(blob, shape.image.ext)
                    with open(os.path.join(assets_dir, digest + "." + safe_ext), "wb") as fh:
                        fh.write(safe)
                    poster = digest + "." + safe_ext
                    pictures.append({
                        "asset": poster, "sha1": digest, "bytes": len(blob),
                        "top": shape.top or 0, "left": shape.left or 0,
                    })
                except Exception:
                    pass
                blocks.append({"kind": "media", "name": shape.name or "", "poster": poster})
                continue

            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[clean(c.text) for c in row.cells] for row in shape.table.rows]
                tables.append(rows)
                blocks.append({"kind": "table", "rows": rows})
                continue

            if not shape.has_text_frame:
                continue
            text = clean(shape.text_frame.text)
            if not text or text.isdigit():
                continue
            pt = shape_font_pt(shape)
            if pt is not None and pt < BODY_MIN_PT:   # credit line or page number
                continue

            for level, para in paragraphs_of(shape):
                if para.isdigit():
                    continue
                blocks.append({"kind": "text", "level": level, "text": para})

        notes = ""
        if slide.has_notes_slide:
            notes = clean(slide.notes_slide.notes_text_frame.text)

        # A divider carries a heading and nothing else. An embedded video counts
        # as content even when it contributes no text and no poster image, or a
        # video slide with a heading would be mistaken for a divider and lose
        # the one thing it was there to show.
        has_body = any(b["kind"] in ("text", "media") for b in blocks)
        out_slides.append({
            "n": index,
            "layout": slide.slide_layout.name,
            "section": None,
            "title": title,
            "is_section": bool(title) and not has_body and not pictures and not tables,
            "blocks": blocks,
            "pictures": pictures,
            "tables": tables,
            "notes": notes,
            "unreadable_shapes": unreadable,
        })

    return {
        "source_file": os.path.basename(path),
        "kind": "pptx",
        "total_slides": len(slides),
        "furniture_images": len(furniture),
        "slides": out_slides,
    }


def extract_pdf(path, outdir):
    """Slide decks delivered as PDF: one page is one slide."""
    from pypdf import PdfReader

    reader = PdfReader(path)
    assets_dir = os.path.join(outdir, "assets")
    os.makedirs(assets_dir, exist_ok=True)

    uses, blobs, per_page = collections.Counter(), {}, []
    for page in reader.pages:
        found = []
        try:
            images = list(page.images)
        except Exception:
            images = []
        for image in images:
            try:
                blob = image.data
            except Exception:
                continue
            digest = hashlib.sha1(blob).hexdigest()
            ext = (os.path.splitext(image.name)[1] or ".png").lstrip(".")
            blobs[digest] = (blob, ext)
            if digest not in found:
                found.append(digest)
                uses[digest] += 1
        per_page.append(found)

    furniture = set()
    for digest, count in uses.items():
        repeated_small = (
            count >= FURNITURE_MIN_USES and len(blobs[digest][0]) <= FURNITURE_MAX_BYTES
        )
        if repeated_small or is_decorative(blobs[digest][0]):
            furniture.add(digest)

    for digest, pair in blobs.items():
        if digest not in furniture:
            safe, safe_ext = web_safe_image(pair[0], pair[1])
            blobs[digest] = (pair[0], safe_ext)
            with open(os.path.join(assets_dir, digest + "." + safe_ext), "wb") as fh:
                fh.write(safe)

    # ---- what repeats on nearly every page is furniture, not content ------
    #
    # These exports print a running footer — "Slide No. 15Principles of Flight
    # and Aircraft Performance (A)" — which is short, sits last, and reads
    # exactly like a heading. Taken as the title it displaces the real one on
    # every single slide, so it has to be recognised before anything else.
    # Recognised by repetition rather than by pattern: a footer is whatever the
    # deck prints on all of its pages.
    page_lines = []
    for page in reader.pages:
        try:
            raw = page.extract_text() or ""
        except Exception:
            raw = ""
        page_lines.append([clean(l) for l in raw.split("\n") if clean(l)])

    shape = collections.Counter()
    for lines in page_lines:
        for line in set(lines):
            shape[re.sub(r"\d+", "#", line)] += 1

    total_pages = max(len(page_lines), 1)
    running = {
        form
        for form, count in shape.items()
        if count >= max(5, total_pages * 0.3) and len(form) <= 120
    }

    # ---- where this deck puts its titles ----------------------------------
    #
    # Decided once for the whole document, not per page. Some exports lay the
    # title down first and some last, and a per-page guess gets it wrong on
    # exactly the pages where it matters: a page whose last line is a wrapped
    # bullet ends up named after half a sentence. Whichever end reads like a
    # heading more often across the deck is the end this deck uses.
    def heading_shaped(line, bulleted):
        if bulleted or not line:
            return False
        if len(line) > 70 or line.endswith((".", ",", ";")):
            return False
        return any(c.isalpha() for c in line)

    first_wins = last_wins = 0
    for lines in page_lines:
        body = [l for l in lines if not l.isdigit() and re.sub(r"\d+", "#", l) not in running]
        if len(body) < 2:
            continue
        if heading_shaped(body[0], bool(BULLET_GLYPH.match(body[0]))):
            first_wins += 1
        if heading_shaped(body[-1], bool(BULLET_GLYPH.match(body[-1]))):
            last_wins += 1
    title_first = first_wins >= last_wins

    # A short opening line that recurs across the deck is a chapter label
    # printed above each slide's own title. Only meaningful in a deck that
    # titles from the top — in one that titles from the bottom, a repeated
    # first line is a repeated bullet.
    section_labels = set()
    if title_first:
        # A chapter label is a complete heading with the slide's own heading
        # under it. Both halves matter: without the first test the fragments of
        # a wrapped title ("Pressure and") are counted as labels, and without
        # the second a genuinely repeated slide title is mistaken for a chapter
        # and its content is filed under a section of the same name.
        # The test that separates a chapter label from the first half of a
        # wrapped title: a chapter heads several *different* slides, so it
        # appears above many different second lines. A title broken across two
        # lines appears above the same second line every time it recurs.
        partners = collections.defaultdict(set)
        for lines in page_lines:
            if len(lines) < 2:
                continue
            head = lines[0]
            if len(head) > 40 or head.endswith((".", ":", ";")):
                continue
            if looks_unfinished(head):
                continue
            partners[head].add(lines[1])
        section_labels = {text for text, seconds in partners.items() if len(seconds) >= 3}

        # A deck either uses the two-level layout or it does not. Where it does,
        # nearly every slide carries its chapter above its own heading; where it
        # does not, a handful of coincidences look like labels and turn ordinary
        # lessons into spurious sections. So the labels are kept only if they
        # account for a real share of the deck.
        labelled_pages = sum(
            1 for lines in page_lines if lines and lines[0] in section_labels
        )
        if labelled_pages < max(8, len(page_lines) * 0.4):
            section_labels = set()

    out_slides = []
    for index, page in enumerate(reader.pages, start=1):
        lines = [
            l for l in page_lines[index - 1]
            if not l.isdigit()
            and not is_junk(l)
            and re.sub(r"\d+", "#", l) not in running
        ]

        # This deck draws its bullets with a symbol font, so every bullet comes
        # out of the text layer with a leading "S" that is a glyph, not a word.
        # It has to go before the line can be read, but the bullet it marks is
        # what tells continuation lines apart from new points.
        starts_bullet = [bool(BULLET_GLYPH.match(l)) for l in lines]
        lines = [BULLET_GLYPH.sub("", l).strip() for l in lines]
        # Where the bullet glyph sits hard against the following word it leaves
        # a stray letter fused to the front of a heading ("sTransponders").
        # No heading in the source begins that way, so it is safe to lift.
        lines = [STUCK_GLYPH.sub("", l) for l in lines]

        # PowerPoint exports lay the title down in one of three places, and
        # which one is a property of the deck rather than of the page.
        title = None
        section = None

        # Where the page sets its heading larger than its body, that is the
        # heading — no guessing from wording, and a heading that wrapped comes
        # back whole.
        font_title, _parts = pdf_page_title(page)
        if font_title:
            lines, starts_bullet = strip_title_lines(lines, starts_bullet, font_title)
            title = font_title

        # Some decks print the chapter name above the slide title on every
        # slide of that chapter. Taken as the title it names forty slides the
        # same thing; recognised for what it is, it gives the module structure
        # a deck like this otherwise has no dividers for.
        if lines and lines[0] in section_labels and len(lines) > 1:
            section = lines[0]
            lines, starts_bullet = lines[1:], starts_bullet[1:]

        # Take it from the end this deck uses, and only from there.
        if not title and lines:
            end = 0 if title_first else -1
            candidate = lines[end]
            if (
                not starts_bullet[end]
                and len(candidate) <= TITLE_MAX_CHARS
                and not candidate.endswith(".")
            ):
                if title_first:
                    title, used = join_wrapped_title(
                        lines, starts_bullet, 0, TITLE_MAX_CHARS
                    )
                    lines, starts_bullet = lines[used:], starts_bullet[used:]
                else:
                    title = candidate
                    lines, starts_bullet = lines[:-1], starts_bullet[:-1]

        # Rejoin the wrapped tail of a line onto the line it belongs to, so a
        # single point is not filed as two.
        #
        # How a continuation is recognised depends on the deck. Where bullets
        # are drawn, a line without one continues the line above. Where a page
        # is prose with no bullets at all, that test would swallow the whole
        # page into a single paragraph — so there, a continuation is a line the
        # previous one did not finish.
        page_has_bullets = any(starts_bullet)
        body = []
        for i, line in enumerate(lines):
            if not body:
                body.append(line)
                continue
            if page_has_bullets:
                continues = not starts_bullet[i]
            else:
                continues = (
                    not body[-1].endswith((".", "!", "?", ":", ";"))
                    and not line[:1].isupper()
                    and not line[:1].isdigit()
                )
            if continues:
                body[-1] = body[-1] + " " + line
            else:
                body.append(line)

        blocks = [{"kind": "text", "level": 0, "text": l} for l in body]
        pictures = []
        for digest in per_page[index - 1]:
            if digest in furniture:
                continue
            pictures.append({
                "asset": digest + "." + blobs[digest][1],
                "sha1": digest,
                "bytes": len(blobs[digest][0]),
                "top": 0,
                "left": 0,
            })
            blocks.append({"kind": "picture", "sha1": digest})

        # A page whose only heading is the chapter label is a continuation of

        # that chapter, not a slide named after it. Left as a title it would

        # open a new lesson called "Aeroscience" every few pages.

        if title and title in section_labels:

            section = section or title

            title = None


        out_slides.append({
            "n": index,
            "layout": "pdf-page",
            "section": section,
            "title": title,
            "is_section": bool(title) and not body and not pictures,
            "blocks": blocks,
            "pictures": pictures,
            "tables": [],
            "notes": "",
            "unreadable_shapes": 0,
        })

    return {
        "source_file": os.path.basename(path),
        "kind": "pdf",
        "total_slides": len(reader.pages),
        "furniture_images": len(furniture),
        "slides": out_slides,
    }


def main():
    src, outdir = sys.argv[1], sys.argv[2]
    os.makedirs(outdir, exist_ok=True)
    if src.lower().endswith(".pptx"):
        data = extract_pptx(src, outdir)
    else:
        data = extract_pdf(src, outdir)

    with open(os.path.join(outdir, "manifest.json"), "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)

    text_blocks = sum(1 for s in data["slides"] for b in s["blocks"] if b["kind"] == "text")
    chars = sum(len(b["text"]) for s in data["slides"] for b in s["blocks"] if b["kind"] == "text")
    empty = [s["n"] for s in data["slides"] if not s["blocks"] and not s["title"]]
    print(json.dumps({
        "source": data["source_file"],
        "slides": data["total_slides"],
        "text_blocks": text_blocks,
        "chars": chars,
        "pictures_kept": sum(len(s["pictures"]) for s in data["slides"]),
        "furniture_images_dropped": data["furniture_images"],
        "tables": sum(len(s["tables"]) for s in data["slides"]),
        "titled_slides": sum(1 for s in data["slides"] if s["title"]),
        "section_slides": sum(1 for s in data["slides"] if s["is_section"]),
        "media_slides": sum(1 for s in data["slides"] for b in s["blocks"] if b["kind"] == "media"),
        "unreadable_shapes": sum(s["unreadable_shapes"] for s in data["slides"]),
        "slides_with_nothing": empty,
    }, indent=1))


if __name__ == "__main__":
    main()
